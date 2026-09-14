"""DOCX export for ATS-safe tailored CVs."""
from __future__ import annotations

import io
import re

from app.schemas.cv import CvTailoredSections


def sanitize_docx_filename(name: str, job_title: str | None = None) -> str:
    base = "-".join(part for part in [name, job_title or "tailored"] if part.strip())
    cleaned = re.sub(r"[^\w.\-]+", "-", base.strip().lower()).strip("-")
    return f"{cleaned or 'tailored-cv'}-cv.docx"


def build_tailored_cv_docx(sections: CvTailoredSections) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    name = sections.contact.name or "Tailored CV"
    title = document.add_paragraph()
    run = title.add_run(name)
    run.bold = True
    run.font.size = Pt(14)

    contact_parts = [
        sections.contact.email,
        sections.contact.phone,
        sections.contact.location,
        sections.contact.linkedin,
    ]
    contact_line = " | ".join(part for part in contact_parts if part)
    if contact_line:
        document.add_paragraph(contact_line)

    def heading(text: str) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(12)

    if sections.summary:
        heading("Professional Summary")
        document.add_paragraph(sections.summary)

    if sections.skills:
        heading("Skills")
        document.add_paragraph(", ".join(sections.skills))

    if sections.experience:
        heading("Experience")
        for item in sections.experience:
            role_line = " | ".join(part for part in [item.title, item.company, item.dates] if part)
            if role_line:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(role_line)
                run.bold = True
            for bullet in item.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if sections.education:
        heading("Education")
        for item in sections.education:
            line = " | ".join(part for part in [item.degree, item.institution, item.dates] if part)
            if line:
                paragraph = document.add_paragraph()
                paragraph.add_run(line).bold = True
            if item.details:
                document.add_paragraph(item.details)

    if sections.certifications:
        heading("Certifications")
        for certification in sections.certifications:
            document.add_paragraph(certification, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
